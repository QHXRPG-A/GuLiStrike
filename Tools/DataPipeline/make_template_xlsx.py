"""Generate Data/Excel/GuLiStrikeShip.xlsx from CDO seed values (one-shot, reusable).

v2 表格式：每个 sheet 前三行是元数据（列名 / 类型 / 必要性），数据从第 4 行起；
标准三列 id/name/Note 必有（name 列 = DataTable 行名）。

Input:  Data/tmp_cdo_read.json  (produced by Scripts/read_ship_cdos.py in the live editor)
Output: Data/Excel/GuLiStrikeShip.xlsx with sheets "Parts" and "Tuning"

注意：会覆盖现有 GuLiStrikeShip.xlsx —— 它是"从蓝图 CDO 重新拉一套初始表"的
引导工具，日常调数值请直接改表，不要重跑本脚本。

Run: python Tools/DataPipeline/make_template_xlsx.py
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[2]
SEED_PATH = PROJECT_ROOT / "Data" / "tmp_cdo_read.json"
XLSX_PATH = PROJECT_ROOT / "Data" / "Excel" / "GuLiStrikeShip.xlsx"

# (列名, 类型, 必要性) —— 与 v2 约定一致；向量是 PrefixX/Y/Z 三列 float
PARTS_COLUMNS = [
    ("id", "int", "Necessary"), ("name", "str", "Necessary"), ("Note", "str", "Optional"),
    ("PartId", "str", "Necessary"), ("Type", "str", "Necessary"),
    ("PartMass", "float", "Optional"), ("Thrust", "float", "Optional"),
    ("Damage", "float", "Optional"), ("FireRate", "float", "Optional"),
    ("MuzzleX", "float", "Optional"), ("MuzzleY", "float", "Optional"),
    ("MuzzleZ", "float", "Optional"), ("ProjectileClass", "softclass", "Optional"),
]

TUNING_COLUMNS = [
    ("id", "int", "Necessary"), ("name", "str", "Necessary"), ("Note", "str", "Optional"),
    ("HullMass", "float", "Necessary"), ("BaseMaxSpeed", "float", "Necessary"),
    ("BaseAcceleration", "float", "Necessary"), ("NominalThrustRatio", "float", "Necessary"),
    ("SpeedMultiplierMin", "float", "Necessary"), ("SpeedMultiplierMax", "float", "Necessary"),
    ("BoostThrustMultiplier", "float", "Necessary"), ("MousePitchScale", "float", "Necessary"),
    ("MouseYawScale", "float", "Necessary"), ("YawRate", "float", "Necessary"),
    ("YawResponseSpeed", "float", "Necessary"), ("YawStopDamping", "float", "Necessary"),
    ("MaxBankAngle", "float", "Necessary"), ("BankInterpSpeed", "float", "Necessary"),
    ("OrientTurnSpeed", "float", "Necessary"), ("OrientToMovement", "bool", "Necessary"),
    ("OrientMinForwardDot", "float", "Necessary"), ("CameraPitchMin", "float", "Necessary"),
    ("CameraPitchMax", "float", "Necessary"),
    ("HullMeshOffsetX", "float", "Necessary"), ("HullMeshOffsetY", "float", "Necessary"),
    ("HullMeshOffsetZ", "float", "Necessary"),
]

META_FILL = PatternFill("solid", fgColor="DDEBF7")
META_FONT = Font(bold=True)


def write_meta_rows(ws, columns):
    """前 3 行元数据：列名 / 类型 / 必要性。"""
    for r_idx, idx in ((1, 0), (2, 1), (3, 2)):
        for c_idx, col in enumerate(columns, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=col[idx])
            cell.fill = META_FILL
            cell.font = META_FONT
            cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A4"
    for c_idx, col in enumerate(columns, start=1):
        width = max(len(col[0]) * 1.6, 12)
        ws.column_dimensions[get_column_letter(c_idx)].width = min(width, 42)


def part_id_from_asset(asset_path: str) -> str:
    """PartId 约定 = 蓝图资产名去掉 BP_ 前缀。"""
    name = asset_path.rsplit("/", 1)[-1]
    return name[3:] if name.startswith("BP_") else name


def build_parts_sheet(ws, parts):
    write_meta_rows(ws, PARTS_COLUMNS)
    next_id = 1
    for p in sorted(parts, key=lambda x: (not x["is_engine"], x["asset_path"])):
        pid = part_id_from_asset(p["asset_path"])
        if p["is_engine"]:
            values = {
                "PartId": pid, "Type": "Engine", "PartMass": p["part_mass"], "Thrust": p["thrust"],
            }
        else:
            mo = p["muzzle_offset"]
            values = {
                "PartId": pid, "Type": "Weapon", "PartMass": p["part_mass"],
                "Damage": p["damage"], "FireRate": round(p["fire_rate"], 4),
                "MuzzleX": mo[0], "MuzzleY": mo[1], "MuzzleZ": mo[2],
                "ProjectileClass": p.get("projectile_class"),
            }
        values["name"] = pid  # 兜底显示名；行名需唯一，用 PartId 保证
        ws.append([next_id if c[0] == "id" else values.get(c[0]) for c in PARTS_COLUMNS])
        next_id += 1


def build_tuning_sheet(ws, tuning):
    write_meta_rows(ws, TUNING_COLUMNS)
    hmo = tuning["hull_mesh_offset"]
    values = {
        "name": "Default",
        "HullMass": tuning["hull_mass"], "BaseMaxSpeed": tuning["base_max_speed"],
        "BaseAcceleration": tuning["base_acceleration"],
        "NominalThrustRatio": tuning["nominal_thrust_ratio"],
        "SpeedMultiplierMin": tuning["speed_multiplier_min"],
        "SpeedMultiplierMax": tuning["speed_multiplier_max"],
        "BoostThrustMultiplier": tuning["boost_thrust_multiplier"],
        "MousePitchScale": tuning["mouse_pitch_scale"], "MouseYawScale": tuning["mouse_yaw_scale"],
        "YawRate": tuning["yaw_rate"], "YawResponseSpeed": tuning["yaw_response_speed"],
        "YawStopDamping": tuning["yaw_stop_damping"], "MaxBankAngle": tuning["max_bank_angle"],
        "BankInterpSpeed": tuning["bank_interp_speed"],
        "OrientTurnSpeed": tuning["orient_turn_speed"],
        "OrientToMovement": tuning["orient_to_movement"],
        "OrientMinForwardDot": tuning["orient_min_forward_dot"],
        "CameraPitchMin": tuning["camera_pitch_min"], "CameraPitchMax": tuning["camera_pitch_max"],
        "HullMeshOffsetX": hmo[0], "HullMeshOffsetY": hmo[1], "HullMeshOffsetZ": hmo[2],
    }
    ws.append([1 if c[0] == "id" else values.get(c[0]) for c in TUNING_COLUMNS])


def main() -> int:
    if not SEED_PATH.exists():
        print(f"seed file missing: {SEED_PATH}\nrun Scripts/read_ship_cdos.py in the editor first")
        return 1
    seed = json.loads(SEED_PATH.read_text(encoding="utf-8"))
    if seed.get("errors"):
        print("seed has errors, aborting:")
        for e in seed["errors"]:
            print(" ", e)
        return 1

    XLSX_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    build_parts_sheet(wb.active, seed["parts"])
    wb.active.title = "Parts"
    build_tuning_sheet(wb.create_sheet("Tuning"), seed["ship_tuning"])
    wb.save(XLSX_PATH)
    print(f"wrote {XLSX_PATH} ({len(seed['parts'])} part rows, 1 tuning preset)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
